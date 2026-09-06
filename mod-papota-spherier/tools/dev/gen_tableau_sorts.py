# -*- coding: utf-8 -*-
r"""Recensement des sorts de classe, pour le tri des sorts améliorables.

Les runes de sort ajoutent un rang au-delà du maximum de Blizzard, jusqu'à trois
par sort. Il faut donc choisir SIX sorts améliorables par classe.

Deux origines sont recensées, et distinguées par une colonne :

  - les sorts **enseignés par un maître de classe**, lus dans `trainer` /
    `trainer_spell` ;
  - les sorts **octroyés par un talent**, lus dans `Talent.dbc` et rattachés à
    leur classe par `TalentTab.dbc` (révision du 2026-08-26).

Les sorts de familier restent hors périmètre : ni maître de classe, ni talent.

Les rangs viennent de `spell_ranks`, les noms français et le drapeau passif de
`Spell.dbc` du serveur. **frFR est la locale d'index 2** (0 enUS, 1 koKR, 2 frFR,
3 deDE…) — se tromper d'index donne de l'allemand.

Sortie : un classeur, une feuille par classe, avec une proposition et une colonne
à remplir.
"""

import os as _os_local, sys as _sys_local
_sys_local.path.insert(0, _os_local.path.dirname(_os_local.path.abspath(__file__)))
from config_local import MYSQL_HOTE, MYSQL_PORT, MYSQL_UTILISATEUR, MYSQL_MDP, BASE_WORLD  # ce qui décrit le poste, hors du dépôt
import io
import os
import struct
import subprocess
import tempfile
from collections import defaultdict

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

MYSQL = r"C:\Program Files\MySQL\MySQL Server 8.4\bin\mysql.exe"
BASE = BASE_WORLD
DBC = r"D:\Serveur WoW\server_hard\bin\RelWithDebInfo\Data\dbc\Spell.dbc"
SORTIE = r"D:\Serveur WoW\outils_spherier\sorts_ameliorables.xlsx"

IDX_NAME, IDX_RANK, IDX_ATTR = 136, 153, 4
LOC_FRFR = 2
SPELL_ATTR0_PASSIVE = 0x00000040

CLASSES = [(1, "Guerrier"), (2, "Paladin"), (3, "Chasseur"), (4, "Voleur"),
           (5, "Prêtre"), (6, "Chevalier de la mort"), (7, "Chaman"),
           (8, "Mage"), (9, "Démoniste"), (11, "Druide")]

# Un sort est proposé comme améliorable s'il est actif et qu'il a au moins ce
# nombre de rangs : en dessous, la courbe de Blizzard est trop courte pour
# qu'on en extrapole trois de plus avec confiance.
RANGS_MINIMUM = 5
PROPOSITIONS = 6            # à retenir par classe


def requete(sql):
    with tempfile.NamedTemporaryFile("w", suffix=".sql", delete=False,
                                     encoding="utf-8") as f:
        f.write(sql)
        chemin = f.name
    try:
        sortie = subprocess.run(
            [MYSQL, "-h" + MYSQL_HOTE, "-u" + MYSQL_UTILISATEUR, "-p" + MYSQL_MDP,
             "--default-character-set=utf8mb4", "-N", "-B", BASE,
             "-e", "source %s" % chemin.replace("\\", "/")],
            capture_output=True, text=True, encoding="utf-8")
        return [l.split("\t") for l in sortie.stdout.splitlines() if l.strip()]
    finally:
        os.unlink(chemin)


# --- Spell.dbc : noms français, sous-titre de rang, drapeau passif -----------
brut = open(DBC, "rb").read()
nrec, nfield, rsize, ssize = struct.unpack_from("<4I", brut, 4)
chaines = brut[20 + nrec * rsize:]


def texte(off):
    if off <= 0 or off >= len(chaines):
        return ""
    return chaines[off:chaines.index(b"\x00", off)].decode("utf-8", "replace")


infos = {}
for i in range(nrec):
    c = struct.unpack_from("<%dI" % nfield, brut, 20 + i * rsize)
    infos[c[0]] = (texte(c[IDX_NAME + LOC_FRFR]) or texte(c[IDX_NAME]),
                   bool(c[IDX_ATTR] & SPELL_ATTR0_PASSIVE))

# --- sorts de maître, regroupés par famille de rangs ------------------------
lignes = requete(
    "SELECT t.Requirement, ts.SpellId, MIN(ts.ReqLevel), "
    "IFNULL(sr.first_spell_id, ts.SpellId), IFNULL(sr.rank, 1) "
    "FROM trainer t JOIN trainer_spell ts ON ts.TrainerId = t.Id "
    "LEFT JOIN spell_ranks sr ON sr.spell_id = ts.SpellId "
    "WHERE t.Type = 0 "
    "GROUP BY t.Requirement, ts.SpellId, sr.first_spell_id, sr.rank;")

par_classe = defaultdict(lambda: defaultdict(dict))
niveaux = defaultdict(lambda: defaultdict(dict))
origine = defaultdict(dict)             # classe -> famille -> "Maître" / "Talent"
for p in lignes:
    if len(p) != 5:
        continue
    classe, spell, niveau, famille, rang = (int(x) for x in p)
    par_classe[classe][famille][rang] = spell
    niveaux[classe][famille][rang] = niveau
    origine[classe][famille] = "Maître"

# --- sorts octroyés par un talent -------------------------------------------
# Talent.dbc : ID(0), TalentTab(1), Row(2), Col(3), SpellRank[9](4-12).
# TalentTab.dbc : ID(0), ..., ClassMask à l'avant-dernier champ utile.
def lire_dbc(chemin):
    b = open(chemin, "rb").read()
    n, nf, rs, _ = struct.unpack_from("<4I", b, 4)
    for i in range(n):
        yield struct.unpack_from("<%dI" % nf, b, 20 + i * rs)


DOSSIER = os.path.dirname(DBC)
# Masque de classe de chaque onglet de talents.
masques = {}
for c in lire_dbc(os.path.join(DOSSIER, "TalentTab.dbc")):
    for champ in c[len(c) - 5:]:
        if 0 < champ < 4096 and bin(champ).count("1") == 1:
            masques[c[0]] = champ
            break

CLASSE_BIT = {classe: 1 << (classe - 1) for classe, _ in CLASSES}

for c in lire_dbc(os.path.join(DOSSIER, "Talent.dbc")):
    masque = masques.get(c[1])
    if not masque:
        continue
    classe = next((cl for cl, bit in CLASSE_BIT.items() if bit == masque), None)
    if not classe:
        continue
    # Les rangs du talent sont autant de rangs du même sort : le dernier non nul
    # est le maximum atteignable en y mettant tous ses points.
    rangs = [s for s in c[4:13] if s]
    if not rangs:
        continue
    famille = rangs[0]
    if famille in par_classe[classe]:
        continue                        # déjà vu par un maître
    for r, spell in enumerate(rangs, start=1):
        par_classe[classe][famille][r] = spell
        niveaux[classe][famille][r] = 0
    origine[classe][famille] = "Talent"

# --- mise en forme ----------------------------------------------------------
TITRE = Font(bold=True, color="FFFFFF", size=11)
FOND_TITRE = PatternFill("solid", fgColor="4F3A6B")
FOND_SAISIE = PatternFill("solid", fgColor="FFF2CC")
FOND_PROPOSE = PatternFill("solid", fgColor="E2EFDA")
BORDURE = Border(*(Side(style="thin", color="BFBFBF"),) * 4)

wb = Workbook()
wb.remove(wb.active)

ws = wb.create_sheet("Mode d'emploi")
for i, ligne in enumerate([
    "Runes de sort — choix des sorts améliorables",
    "",
    "Une rune de sort ajoute UN rang au-delà du maximum de Blizzard. Elles se cumulent :",
    "deux runes de Pourfendre donnent les rangs 11 et 12, et le joueur connaît les deux.",
    "Retirer une rune fait oublier le rang du dessus. Trois runes au maximum par sort,",
    "donc trois rangs customs à inventer pour chacun.",
    "",
    "Il faut SIX sorts améliorables par classe, soit 60 sorts et 180 rangs à créer.",
    "",
    "Deux origines, distinguées par une colonne : les sorts appris chez un MAÎTRE",
    "de classe, et ceux octroyés par un TALENT. Les sorts de familier restent hors",
    "périmètre.",
    "",
    "Un sort de talent perdu — changement de spécialisation, remise à zéro — désactive",
    "ses runes : les rangs supplémentaires sont oubliés sur-le-champ, sans que le joueur",
    "ait à ouvrir quoi que ce soit. Les runes restent serties et reprennent leur effet",
    "dès que le talent revient.",
    "",
    "Colonne « Proposition » : mon avis. Sont proposés les sorts ACTIFS ayant au moins",
    "cinq rangs — en dessous, la courbe de Blizzard est trop courte pour en extrapoler",
    "trois de plus avec confiance. Les sorts passifs (maîtrises d'armes, armures,",
    "spécialisations) sont écartés d'office : un rang de plus n'y voudrait rien dire.",
    "",
    "Colonne « Améliorable » : votre décision. Six « Oui » par feuille.",
    "Les sorts à rang unique sont conservés dans la liste : vous aviez demandé qu'on",
    "leur invente une courbe, à traiter en dernier en s'appuyant sur la tendance des",
    "autres. Ils restent donc choisissables.",
], start=1):
    c = ws.cell(row=i, column=1, value=ligne)
    if i == 1:
        c.font = Font(bold=True, size=13, color="4F3A6B")
ws.column_dimensions["A"].width = 100

ENTETES = ["Sort", "Origine", "Rangs", "Niveau du dernier rang", "Type",
           "Sort du dernier rang", "Proposition", "Améliorable"]
LARGEURS = [30, 10, 8, 12, 10, 12, 13, 13]

for classe, nom in CLASSES:
    ws = wb.create_sheet(nom)
    familles = []
    for famille, rangs in par_classe[classe].items():
        nom_sort, passif = infos.get(famille, ("?", False))
        maxi = max(rangs)
        familles.append({
            "nom": nom_sort, "rangs": maxi, "passif": passif,
            "niveau": max(niveaux[classe][famille].values()),
            "spell": rangs[maxi],
            "origine": origine[classe].get(famille, "Maître"),
        })

    # Actifs d'abord, du plus de rangs au moins ; passifs relégués à la fin.
    familles.sort(key=lambda d: (d["passif"], -d["rangs"], d["nom"]))
    proposes = 0
    for d in familles:
        d["propose"] = ""
        if (not d["passif"] and d["rangs"] >= RANGS_MINIMUM
                and proposes < PROPOSITIONS):
            d["propose"] = "Oui"
            proposes += 1

    for i, titre in enumerate(ENTETES, start=1):
        c = ws.cell(row=1, column=i, value=titre)
        c.font, c.fill, c.border = TITRE, FOND_TITRE, BORDURE
        c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    ws.row_dimensions[1].height = 30

    for r, d in enumerate(familles, start=2):
        valeurs = [d["nom"], d["origine"], d["rangs"],
                   d["niveau"] or "", "Passif" if d["passif"] else "Actif",
                   d["spell"], d["propose"], None]
        for i, v in enumerate(valeurs, start=1):
            c = ws.cell(row=r, column=i, value=v)
            c.border = BORDURE
            c.alignment = Alignment(vertical="center",
                                    horizontal="left" if i == 1 else "center")
        if d["propose"]:
            ws.cell(row=r, column=7).fill = FOND_PROPOSE
        ws.cell(row=r, column=8).fill = FOND_SAISIE

    dv = DataValidation(type="list", formula1='"Oui,Non"', allow_blank=True,
                        showDropDown=False)
    ws.add_data_validation(dv)
    dv.add("H2:H%d" % (len(familles) + 1))

    for i, largeur in enumerate(LARGEURS, start=1):
        ws.column_dimensions[get_column_letter(i)].width = largeur
    ws.freeze_panes = "A2"

wb.save(SORTIE)
print("Classeur écrit :", SORTIE)
for classe, nom in CLASSES:
    total = len(par_classe[classe])
    actifs = sum(1 for f, r in par_classe[classe].items()
                 if not infos.get(f, ("", False))[1] and max(r) >= RANGS_MINIMUM)
    print("  %-22s %2d familles, dont %2d actives à %d rangs ou plus"
          % (nom, total, actifs, RANGS_MINIMUM))
