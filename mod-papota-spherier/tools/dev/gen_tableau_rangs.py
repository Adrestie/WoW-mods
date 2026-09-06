# -*- coding: utf-8 -*-
r"""Table de relecture des rangs extrapolés, avant toute écriture en DBC.

Une ligne par sort : sa croissance mesurée, la valeur de son dernier rang chez
Blizzard, et celle des trois rangs créés. De quoi juger d'un coup d'œil si la
courbe reste plausible, sort par sort.
"""

import io
import json

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

SOURCE = r"D:\Serveur WoW\outils_spherier\rangs_extrapoles.json"
SORTIE = r"D:\Serveur WoW\outils_spherier\rangs_proposes.xlsx"

TITRE = Font(bold=True, color="FFFFFF", size=11)
FOND_TITRE = PatternFill("solid", fgColor="4F3A6B")
FOND_CREE = PatternFill("solid", fgColor="E2EFDA")
FOND_ALERTE = PatternFill("solid", fgColor="FCE4D6")
BORDURE = Border(*(Side(style="thin", color="BFBFBF"),) * 4)

ENTETES = ["Sort", "Origine", "Rangs existants", "Croissance", "Source du taux",
           "Dernier rang Blizzard", "Rang +1", "Rang +2", "Rang +3"]
LARGEURS = [30, 10, 10, 11, 17, 15, 12, 12, 12]

donnees = json.load(io.open(SOURCE, encoding="utf-8"))

wb = Workbook()
wb.remove(wb.active)

ws = wb.create_sheet("Méthode")
for i, ligne in enumerate([
    "Rangs supplémentaires — proposition",
    "",
    "Mesure faite sur les 1095 rangs existants des 165 sorts retenus : la croissance",
    "d'un rang au suivant, en FIN de courbe, a pour médiane ×1,211, soit +21 % par rang.",
    "Les quartiles vont de ×1,144 à ×1,296.",
    "",
    "Chaque sort garde SA propre croissance quand il a de quoi la mesurer — la médiane",
    "de ses trois derniers pas, et non leur moyenne : une refonte d'extension laisse",
    "parfois un saut isolé qui fausserait une moyenne. Le taux est borné à [1,10 ; 1,35] ;",
    "au-delà, on quitterait ce que fait Blizzard.",
    "",
    "Les sorts trop courts pour être mesurés reçoivent la médiane générale. C'est",
    "exactement l'usage prévu : ils passent en dernier et profitent de la tendance",
    "des autres. 115 sorts ont leur croissance propre, 50 suivent la tendance.",
    "",
    "QUEL EFFET GRANDIT ? Pas de liste de types : un effet est mis à l'échelle s'il a",
    "effectivement varié d'un rang à l'autre chez Blizzard. La Frappe de givre porte",
    "ses dégâts plats dans un effet et un pourcentage d'arme dans un autre — le premier",
    "monte, le second reste. La règle est mesurée, pas décrétée.",
    "",
    "Ce qui ne bouge pas : le niveau requis, qui reste celui du dernier rang Blizzard",
    "(ces rangs s'obtiennent par une rune, et un niveau au-delà de 80 les rendrait",
    "incastables), et le coût, qui chez Blizzard ne varie pas d'un rang à l'autre.",
    "",
    "La colonne « valeur » est le plus gros effet chiffré du rang : dégâts, soins ou",
    "montant d'aura selon le sort. Elle sert à comparer, pas à décrire le sort entier.",
], start=1):
    c = ws.cell(row=i, column=1, value=ligne)
    if i == 1:
        c.font = Font(bold=True, size=13, color="4F3A6B")
ws.column_dimensions["A"].width = 92

for classe, liste in donnees.items():
    ws = wb.create_sheet(classe[:31])
    for i, titre in enumerate(ENTETES, start=1):
        c = ws.cell(row=1, column=i, value=titre)
        c.font, c.fill, c.border = TITRE, FOND_TITRE, BORDURE
        c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    ws.row_dimensions[1].height = 30

    liste = sorted(liste, key=lambda d: (-d["rangs_existants"], d["nom"]))
    for r, d in enumerate(liste, start=2):
        valeurs = [d["nom"], d["origine"], d["rangs_existants"],
                   "×%.3f" % d["taux"], d["origine_taux"],
                   round(d["ampleur_derniere"], 1)]
        valeurs += [round(x["ampleur"], 1) for x in d["rangs"]]
        for i, v in enumerate(valeurs, start=1):
            c = ws.cell(row=r, column=i, value=v)
            c.border = BORDURE
            c.alignment = Alignment(vertical="center",
                                    horizontal="left" if i == 1 else "center")
        for i in (7, 8, 9):
            ws.cell(row=r, column=i).fill = FOND_CREE
        # Un sort dont on n'a rien pu mesurer mérite un coup d'œil.
        if d["ampleur_derniere"] == 0:
            for i in range(1, len(ENTETES) + 1):
                ws.cell(row=r, column=i).fill = FOND_ALERTE

    for i, largeur in enumerate(LARGEURS, start=1):
        ws.column_dimensions[get_column_letter(i)].width = largeur
    ws.freeze_panes = "A2"

wb.save(SORTIE)
total = sum(len(v) for v in donnees.values())
muets = sum(1 for v in donnees.values() for d in v if d["ampleur_derniere"] == 0)
print("Classeur écrit :", SORTIE)
print("  %d sorts, %d rangs proposés ; %d sorts sans valeur mesurable (surlignés)"
      % (total, total * 3, muets))
